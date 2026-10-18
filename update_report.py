import os
import re
import fitz
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import RGBColor

def update_pdf(input_pdf, output_pdf):
    doc = fitz.open(input_pdf)
    
    def adjust_stream(s_text):
        # 1. Colors to pure Black and White
        # Convert any background fill to pure white (1 1 1 rg)
        s_text = re.sub(r'0\.965\s+0\.976\s+0\.984\s+rg', '1 1 1 rg', s_text)
        s_text = re.sub(r'0\.929\s+0\.949\s+0\.965\s+rg', '1 1 1 rg', s_text)
        s_text = re.sub(r'0\.969\s+0\.976\s+0\.984\s+rg', '1 1 1 rg', s_text)
        s_text = re.sub(r'0\.9[0-9]+\s+0\.9[0-9]+\s+0\.9[0-9]+\s+rg', '1 1 1 rg', s_text)
        
        # Convert colored text to pure black
        s_text = re.sub(r'0\.176\s+0\.212\s+0\.282\s+(rg|RG)', r'0 0 0 \1', s_text)
        s_text = re.sub(r'0\s+0\.6\s+0\s+(rg|RG)', r'0 0 0 \1', s_text)
        s_text = re.sub(r'0\.8\s+0\s+0\s+(rg|RG)', r'0 0 0 \1', s_text)
        s_text = re.sub(r'0\.102\s+0\.122\s+0\.169\s+(rg|RG)', r'0 0 0 \1', s_text)
        s_text = re.sub(r'0\.102\s+0\.212\s+0\.365\s+(rg|RG)', r'0 0 0 \1', s_text)
        s_text = re.sub(r'0\.77\s+0\.19\s+0\.19\s+(rg|RG)', r'0 0 0 \1', s_text)

        # 2. Convert thick borders (0.96 pt) and standard borders (0.48 pt) to light normal borders (0.35 pt)
        def replace_re(match):
            x, y, w, h = match.group(1), match.group(2), match.group(3), match.group(4)
            wf, hf = float(w), float(h)
            # Corner / intersection point
            if wf <= 1.05 and hf <= 1.05:
                return f'{x} {y} 0.35 0.35 re'
            # Vertical border
            if wf <= 1.05:
                return f'{x} {y} 0.35 {h} re'
            # Horizontal border
            if hf <= 1.05:
                return f'{x} {y} {w} 0.35 re'
            return match.group(0)

        s_text = re.sub(r'([\d\.]+)\s+([\d\.]+)\s+([\d\.]+)\s+([\d\.]+)\s+re', replace_re, s_text)
        return s_text

    for pno, page in enumerate(doc):
        contents_xref = page.get_contents()
        if contents_xref:
            stream = page.read_contents()
            s_text = stream.decode('latin1')
            s_new = adjust_stream(s_text)
            doc.update_stream(contents_xref[0], s_new.encode('latin1'))

    doc.save(output_pdf)
    print(f"Successfully processed PDF: {output_pdf}")

def update_docx(input_docx, output_docx):
    doc = docx.Document(input_docx)
    ns = nsdecls('w')

    for table in doc.tables:
        tblPr = table._tbl.tblPr
        old_borders = tblPr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders')
        for b in old_borders:
            tblPr.remove(b)
        
        # Set light normal borders: sz="4" is 1/2 pt (0.5 pt) standard light border
        new_borders = parse_xml(
            f'<w:tblBorders {ns}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(new_borders)

        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                # Remove any background shading
                for s in tcPr.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd'):
                    tcPr.remove(s)
                # Ensure all text in cell is pure black
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0, 0, 0)

    # Ensure all document paragraphs and runs are pure black
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)

    # Ensure section page borders are light normal borders (sz="4" is 1/2 pt)
    for section in doc.sections:
        sectPr = section._sectPr
        pgBorders = sectPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgBorders')
        if pgBorders is not None:
            sectPr.remove(pgBorders)
            new_pg_borders = parse_xml(
                f'<w:pgBorders {ns} w:offsetFrom="page">'
                f'<w:top w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
                f'<w:left w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
                f'<w:bottom w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
                f'<w:right w:val="single" w:sz="4" w:space="24" w:color="000000"/>'
                f'</w:pgBorders>'
            )
            sectPr.append(new_pg_borders)

    doc.save(output_docx)
    print(f"Successfully processed DOCX: {output_docx}")

if __name__ == "__main__":
    # 1. Update the uploaded report cv.pdf in place (with backup)
    if os.path.exists("report cv.pdf"):
        if not os.path.exists("report cv_backup.pdf"):
            import shutil
            shutil.copy("report cv.pdf", "report cv_backup.pdf")
        update_pdf("report cv_backup.pdf", "report cv.pdf")
        print("Updated report cv.pdf in project directory.")

    # 2. Update DOCX files
    downloads_docx = r"C:\Users\HP\Downloads\Safety_Helmet_Detection_Project_Report.docx"
    if os.path.exists(downloads_docx):
        update_docx(downloads_docx, "Safety_Helmet_Detection_Project_Report.docx")
        update_docx(downloads_docx, "report cv.docx")
        print("Updated Safety_Helmet_Detection_Project_Report.docx and report cv.docx.")
